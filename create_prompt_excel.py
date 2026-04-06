import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

wb = openpyxl.Workbook()

# ─────────────────────────────────────────
# 共通スタイル
# ─────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=12)
LABEL_FILL  = PatternFill("solid", fgColor="D6E4F0")
LABEL_FONT  = Font(bold=True, size=11)
INPUT_FILL  = PatternFill("solid", fgColor="FFFEF0")   # 薄黄色（入力欄）
INPUT_FONT  = Font(size=11)
INQUIRY_FILL = PatternFill("solid", fgColor="F0FFF0")  # 薄緑（問合せ欄）
THIN   = Side(style="thin",   color="AAAAAA")
MEDIUM = Side(style="medium", color="1F4E79")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
TOP    = Alignment(horizontal="left",   vertical="top",    wrap_text=True)

def cell(ws, row, col, value="", font=None, fill=None, align=None, border=None):
    c = ws.cell(row=row, column=col, value=value)
    if font:   c.font      = font
    if fill:   c.fill      = fill
    if align:  c.alignment = align
    if border: c.border    = border
    return c

# ─────────────────────────────────────────
# Sheet 1: 入力シート
# ─────────────────────────────────────────
ws1 = wb.active
ws1.title = "シート1_入力"

# タイトル
ws1.merge_cells("A1:C1")
cell(ws1, 1, 1, "プロンプト生成ツール ／ 入力シート",
     font=Font(bold=True, color="FFFFFF", size=14),
     fill=HEADER_FILL, align=CENTER)

# 説明
ws1.merge_cells("A2:C2")
cell(ws1, 2, 1,
     "4項目を入力してください。シート2にプロンプトが自動生成されます。"
     "　※ 問合せ事項はメール文をそのままペーストしてもOKです。",
     font=Font(size=10, color="444444", italic=True), align=LEFT)

# ヘッダー行
for col, text in enumerate(["No.", "項目名", "入力値"], start=1):
    cell(ws1, 3, col, text, font=HEADER_FONT, fill=HEADER_FILL,
         align=CENTER, border=BORDER)

# ── 行1: 案件名 ──
cell(ws1, 4, 1, "1", font=LABEL_FONT, fill=LABEL_FILL, align=CENTER, border=BORDER)
cell(ws1, 4, 2, "案件名", font=LABEL_FONT, fill=LABEL_FILL, align=LEFT, border=BORDER)
c = ws1.cell(row=4, column=3, value="（例）育児休業申請対応")
c.font = INPUT_FONT; c.fill = INPUT_FILL; c.alignment = LEFT; c.border = BORDER

# ── 行2: 担当者 ──
cell(ws1, 5, 1, "2", font=LABEL_FONT, fill=LABEL_FILL, align=CENTER, border=BORDER)
cell(ws1, 5, 2, "担当者（総務/人事/経理）\nおよび決裁者（社長/役員）",
     font=LABEL_FONT, fill=LABEL_FILL, align=LEFT, border=BORDER)
c = ws1.cell(row=5, column=3, value="総務担当者および社長")
c.font = INPUT_FONT; c.fill = INPUT_FILL; c.alignment = LEFT; c.border = BORDER

# ── 行3: 成果物タイプ ──
cell(ws1, 6, 1, "3", font=LABEL_FONT, fill=LABEL_FILL, align=CENTER, border=BORDER)
cell(ws1, 6, 2, "成果物タイプ\n（ドロップダウンから選択）",
     font=LABEL_FONT, fill=LABEL_FILL, align=LEFT, border=BORDER)
c = ws1.cell(row=6, column=3, value="回答資料")
c.font = INPUT_FONT; c.fill = INPUT_FILL; c.alignment = LEFT; c.border = BORDER

dv = DataValidation(
    type="list",
    formula1='"回答資料,返信メール文,プレゼン"',
    allow_blank=False, showDropDown=False,
    showErrorMessage=True,
    errorTitle="入力エラー",
    error="「回答資料」「返信メール文」「プレゼン」の中から選択してください。"
)
ws1.add_data_validation(dv)
dv.add(ws1["C6"])

# ── 行4: 問合せ事項（大きめ入力欄）──
cell(ws1, 7, 1, "4", font=LABEL_FONT, fill=LABEL_FILL, align=CENTER, border=BORDER)
cell(ws1, 7, 2,
     "問合せ事項\n（メール文のペースト可）",
     font=LABEL_FONT, fill=LABEL_FILL, align=LEFT, border=BORDER)
c = ws1.cell(row=7, column=3,
             value="（例）先日の社員Aの育児休業取得について、"
                   "いつから申請できるか、手続きの流れを教えてください。")
c.font = INPUT_FONT
c.fill = INQUIRY_FILL
c.alignment = TOP
c.border = Border(
    left=Side(style="medium", color="2E75B6"),
    right=Side(style="medium", color="2E75B6"),
    top=Side(style="medium", color="2E75B6"),
    bottom=Side(style="medium", color="2E75B6"),
)

# 列幅・行高
ws1.column_dimensions["A"].width = 6
ws1.column_dimensions["B"].width = 30
ws1.column_dimensions["C"].width = 60
ws1.row_dimensions[1].height = 30
ws1.row_dimensions[2].height = 22
ws1.row_dimensions[3].height = 22
for r in [4, 5, 6]:
    ws1.row_dimensions[r].height = 36
ws1.row_dimensions[7].height = 120   # 問合せ欄は広く

# ─────────────────────────────────────────
# Sheet 2: プロンプト生成シート
# ─────────────────────────────────────────
ws2 = wb.create_sheet("シート2_プロンプト")

ws2.merge_cells("A1:B1")
cell(ws2, 1, 1, "生成プロンプト（コピーしてご利用ください）",
     font=Font(bold=True, color="FFFFFF", size=13),
     fill=HEADER_FILL, align=CENTER)

# ── プロンプト数式 ──
# シート1参照: C4=案件名, C5=担当者, C6=成果物タイプ, C7=問合せ事項
S = "'シート1_入力'!"

formula = (
    '="[#変数設定]"&CHAR(10)'
    f'&"- 案件名: "&{S}C4&CHAR(10)'
    f'&"- 担当者・決裁者: "&{S}C5&CHAR(10)'
    f'&"- 成果物タイプ: "&{S}C6&CHAR(10)'
    '&"- 問合せ事項:"&CHAR(10)'
    f'&{S}C7&CHAR(10)'
    '&CHAR(10)'
    '&"---"&CHAR(10)'
    '&CHAR(10)'
    '&"# 前提条件"&CHAR(10)'
    f'&"- タイトル: "&{S}C4&"顧問先問い合わせ対応"&CHAR(10)'
    '&"- 依頼者条件: 給与計算代行・労務顧問を行う社会保険労務士法人の実務担当者"&CHAR(10)'
    '&"- 目的と目標:"&CHAR(10)'
    '&"  目的: 顧問先の意思決定と実務処理が止まらないよう、正確・簡潔・根拠付きで回答する"&CHAR(10)'
    '&"  目標:"&CHAR(10)'
    '&"  1) 結論先出し（3行以内）＋理由＋手続/対応手順まで一気通貫"&CHAR(10)'
    '&"  2) 期限・提出先・必要書類・リスクを明示"&CHAR(10)'
    '&"  3) 不足情報がある場合は、推測で断定せず確認質問に切り替える"&CHAR(10)'
    '&CHAR(10)'
    '&"- 前提知識:"&CHAR(10)'
    '&"  ソース内の人事労務・法令・制度要件を前提"&CHAR(10)'
    '&"- 守秘・個人情報:"&CHAR(10)'
    '&"  個人名/マイナンバー/住所等の個人情報は出力しない。出す必要がある場合は匿名化する（例：A氏、従業員1）"&CHAR(10)'
    '&CHAR(10)'
    '&"# ペルソナ設定"&CHAR(10)'
    '&"- 書き手: 社会保険労務士法人の実務責任者。事実と根拠を重視し、顧問先の手間を減らす書き方をする。"&CHAR(10)'
    f'&"- 読み手: 顧問先の"&{S}C5&"。忙しく、結論と次のアクションが欲しい。"&CHAR(10)'
    f'&"- 成果物タイプ: "&{S}C6&CHAR(10)'
    '&CHAR(10)'
    '&"# 実行指示"&CHAR(10)'
    f'&"[#変数設定]・[#参考情報]に基づき、"&{S}C6&" を作成してください。"&CHAR(10)'
    f'&"成果物タイプは "&{S}C6&" に従い、該当する[#出力フォーマット]を厳守してください。"&CHAR(10)'
    '&CHAR(10)'
    '&"## 0. 重要ルール（品質と安全）"&CHAR(10)'
    '&"- まず「前提が足りるか」を判定する。"&CHAR(10)'
    '&"- 足りない場合：[#確認質問]を最大5つ出し、その後に「現時点で言える暫定回答（断定禁止）」と「次アクション」を書く。"&CHAR(10)'
    '&"- 足りる場合：確認質問は出さず、最終成果物を作る。"&CHAR(10)'
    '&"- 断定条件:"&CHAR(10)'
    '&"  - 法令・制度・手続きの""一般論""は断定してよい（根拠URLを併記）。"&CHAR(10)'
    '&"  - 個別事案の適法/違法の断定、裁判例の結論断定、税務判断は避ける。必要なら「要個別検討」とし、追加ヒアリング項目を提示する。"&CHAR(10)'
    '&"- 表現:"&CHAR(10)'
    '&"  - 日本語で簡潔に。専門用語は（ ）で補足。"&CHAR(10)'
    '&"  - 不確実な点は「不明」「要確認」「運用差あり」と明記する。"&CHAR(10)'
    '&"- 計算:"&CHAR(10)'
    '&"  - 数値が必要で入力がない場合、仮定計算は「仮定」と明記し、必要な入力値を列挙する。"&CHAR(10)'
    '&CHAR(10)'
    '&"## 1. 問い合わせを分類する（必須）"&CHAR(10)'
    '&"{問い合わせ分類候補} のどれかに分類し、分類理由を1行で説明する。"&CHAR(10)'
    '&CHAR(10)'
    '&"## 2. 事実関係の整理（必須）"&CHAR(10)'
    '&"- 時系列（いつ/誰/何が起きた）"&CHAR(10)'
    '&"- 雇用形態・就業形態・賃金の締日支払日・規程有無"&CHAR(10)'
    '&"- 争点（質問者が本当に知りたいこと）を1文で定義"&CHAR(10)'
    '&CHAR(10)'
    '&"## 3. 結論（必須）"&CHAR(10)'
    '&"読み手が最初に読む「結論」を3行以内で提示する。"&CHAR(10)'
    '&CHAR(10)'
    '&"## 4. 根拠（必須）"&CHAR(10)'
    '&"- 関連法令/手続/公的機関の案内を、要点→リンクの順で示す（リンクは[#参考情報]優先）"&CHAR(10)'
    '&"- 根拠が複数ある場合、優先順位を付ける（一次情報＞解説記事＞実務慣行）"&CHAR(10)'
    '&CHAR(10)'
    '&"## 5. 実務の手順に落とす（必須）"&CHAR(10)'
    '&"- 顧問先がやること（チェックリスト）"&CHAR(10)'
    '&"- 当法人がやること（受託範囲）"&CHAR(10)'
    '&"- 必要書類/添付資料/提出先/期限（わかる範囲で）"&CHAR(10)'
    '&"- リスクと代替案（選択肢がある場合）"&CHAR(10)'
    '&CHAR(10)'
    '&"## 6. 追加確認が必要な場合の分岐（必須）"&CHAR(10)'
    '&"不足情報がある場合は、確認質問→暫定回答→次アクションの順で書く。"&CHAR(10)'
    '&CHAR(10)'
    '&"## 7. 体裁（必須）"&CHAR(10)'
    '&"- 全体の構成は、以下の[#出力フォーマット]を厳守する。"&CHAR(10)'
    '&"- 指示の復唱はしない。"&CHAR(10)'
    '&"- 前置き・結びの挨拶は不要（ただし、確認質問ブロックは例外として許可）。"&CHAR(10)'
    '&CHAR(10)'
    '&"# 出力フォーマット"&CHAR(10)'
    f'&{S}C6&" に応じて、以下のいずれか「1つだけ」を採用し、他形式は混ぜない。"'
)

ws2.merge_cells("A2:B2")
cp = ws2["A2"]
cp.value = formula
cp.font  = Font(size=10, name="Meiryo")
cp.fill  = PatternFill("solid", fgColor="F2F9FF")
cp.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
cp.border = Border(left=MEDIUM, right=MEDIUM, top=MEDIUM, bottom=MEDIUM)

ws2.column_dimensions["A"].width = 110
ws2.column_dimensions["B"].width = 2
ws2.row_dimensions[1].height = 28
ws2.row_dimensions[2].height = 750

ws2.merge_cells("A3:B3")
cell(ws2, 3, 1,
     "▲ 上のセルをクリック → Ctrl+C でプロンプト全文をコピーできます",
     font=Font(size=9, italic=True, color="666666"),
     align=Alignment(horizontal="left", vertical="center"))

# 保存
out_path = "/home/user/story20250803/プロンプト生成ツール.xlsx"
wb.save(out_path)
print(f"保存完了: {out_path}")
